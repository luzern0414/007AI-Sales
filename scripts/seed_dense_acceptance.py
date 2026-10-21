"""Create isolated compact-layout fixtures through HTTP; no production data."""
import io,json,sys,wave
from pathlib import Path
import httpx
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
with httpx.Client(base_url='http://127.0.0.1:18767',timeout=120) as c:
    c.post('/api/auth/login',json={'username':'workflow-ui','password':'WorkflowUI!2026'}).raise_for_status()
    projects=c.get('/api/projects').json()
    p=next((p for p in projects if p['name']=='紧凑验收 · 协同审批'),None)
    if not p:
        p=c.post('/api/projects',json={'name':'紧凑验收 · 协同审批','customer':'验收客户'}).json()
        base='/api/projects/'+p['id']
        d=Document();d.add_heading('审批方案',0);d.add_paragraph('按部门隔离审批权限，接口仍待开放。');buf=io.BytesIO();d.save(buf)
        files=[('审批方案.docx',buf.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
               ('客户提供-需求.md','# 需求确认\n\n**审批权限**按部门隔离。\n\n|需求|状态|\n|---|---|\n|接口|待确认|'.encode(),'text/markdown'),
               ('数据库设计.md','# 数据库设计\n\n审批表和权限表设计中。'.encode(),'text/markdown')]
        pdf=next((ROOT/'output/playwright').glob('*.pdf'))
        files.append(('审批方案.pdf',pdf.read_bytes(),'application/pdf'))
        for name,payload,mime in files:c.post(base+'/files',files={'file':(name,payload,mime)}).raise_for_status()
        buf=io.BytesIO()
        with wave.open(buf,'wb') as w:w.setnchannels(1);w.setsampwidth(2);w.setframerate(16000);w.writeframes(b'\x00\x00'*4000)
        c.post(base+'/files',files={'file':('20261013-143000-客户会议.wav',buf.getvalue(),'audio/wav')}).raise_for_status()
        m=c.get(base).json()['meetings'][0]
        c.patch('/api/meetings/'+m['id'],json={'version':m['version'],'transcript':'客户接口尚未开放，需求确认后再提交方案。'}).raise_for_status()
        c.post(base+'/files',files={'file':('20261013-143000-客户会议-会议纪要.md','# 客户会议\n\n**接口开放时间**待确认。'.encode(),'text/markdown')}).raise_for_status()
        c.post(base+'/analyze',json={'background':'验收环境：客户使用自研系统。','instructions':'关注接口及时间风险。'}).raise_for_status()
    (ROOT/'output/playwright/workflow/dense-project.json').write_text(json.dumps({'id':p['id']},ensure_ascii=False),encoding='utf-8')
    print('Isolated compact project:',p['id'])
