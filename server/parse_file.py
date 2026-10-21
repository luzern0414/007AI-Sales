"""Bounded, local business-document extraction; never executes active content."""
import base64
import json
import posixpath
import re
import sys
import urllib.parse
import zipfile
import zlib
from html.parser import HTMLParser
from pathlib import Path
from xml.etree import ElementTree as ET

MAX_CHARS = 1_000_000
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_PDF_BYTES = 100 * 1024 * 1024
MAX_ZIP_BYTES = 100 * 1024 * 1024
MAX_XML_BYTES = 20 * 1024 * 1024
MAX_PDF_STREAM_BYTES = 20 * 1024 * 1024
MAX_PAGES = 300
IMAGE_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.webp', '.gif', '.bmp', '.tif', '.tiff'}
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


class _Text:
    def __init__(self):
        self.parts = []
        self.count = 0

    def append(self, value):
        # Include joining newlines in the bound, and never return partial output.
        self.count += len(value) + bool(self.parts)
        if self.count > MAX_CHARS:
            raise ValueError('document too long')
        self.parts.append(value)

    def text(self):
        return '\n'.join(self.parts)


class _HTMLText(HTMLParser):
    """Text only: HTMLParser does not fetch URLs or run JavaScript."""
    IGNORED = {'head', 'script', 'style', 'noscript', 'template'}
    BLOCKS = {'p', 'div', 'section', 'article', 'header', 'footer', 'main', 'aside',
              'li', 'ul', 'ol', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'table', 'tr', 'br', 'hr'}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.output = []
        self.count = 0
        self.ignored = []
        self.in_cell = False
        self.row_cell_count = 0

    def add(self, value):
        self.count += len(value)
        if self.count > MAX_CHARS:
            raise ValueError('HTML too long')
        self.output.append(value)

    def handle_starttag(self, tag, attrs):
        if self.ignored:
            if tag == self.ignored[-1]:
                self.ignored.append(tag)
            return
        if tag in self.IGNORED:
            self.ignored.append(tag)
        elif tag in {'td', 'th'}:
            if self.row_cell_count:
                self.add(' | ')
            self.row_cell_count += 1
            self.in_cell = True
        elif tag in self.BLOCKS:
            if tag == 'tr':
                self.row_cell_count = 0
            if not self.in_cell:
                self.add('\n')
            elif tag == 'br':
                self.add(' ')

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        self.handle_endtag(tag)

    def handle_endtag(self, tag):
        if self.ignored:
            if tag == self.ignored[-1]:
                self.ignored.pop()
            return
        if tag in {'td', 'th'}:
            self.in_cell = False
        elif tag in self.BLOCKS and not self.in_cell:
            self.add('\n')

    def handle_data(self, data):
        if not self.ignored:
            self.add(data)

    def text(self):
        return '\n'.join(line for raw in ''.join(self.output).splitlines()
                         if (line := ' '.join(raw.split())))


def _html_text(value):
    # Plain HTML doctypes are safe; entity declarations are unnecessary here.
    if '<!ENTITY' in value.upper():
        raise ValueError('entity declarations are unsupported')
    parser = _HTMLText()
    parser.feed(value)
    parser.close()
    return parser.text()


def _xml(value):
    if len(value) > MAX_XML_BYTES:
        raise ValueError('XML too large')
    source = value.decode('utf-8-sig') if isinstance(value, bytes) else value
    if '<!DOCTYPE' in source.upper() or '<!ENTITY' in source.upper():
        raise ValueError('DTD and entities are unsupported')
    return ET.fromstring(source)


def _check_archive(archive):
    items = archive.infolist()
    if (len(items) > 20000 or len({i.filename for i in items}) != len(items)
            or sum(i.file_size for i in items) > MAX_ZIP_BYTES
            or any(i.file_size > MAX_XML_BYTES or i.flag_bits & 1 for i in items)):
        raise ValueError('archive expansion limit exceeded')


def _read_xml(archive, name):
    with archive.open(name) as stream:
        value = stream.read(MAX_XML_BYTES + 1)
    return _xml(value)


def _drawing_paragraph(paragraph):
    return ''.join(node.text or '' if node.tag == f'{{{A}}}t' else '\n'
                   for node in paragraph.iter()
                   if node.tag in {f'{{{A}}}t', f'{{{A}}}br'})


def _pptx(path, output):
    found = False
    with zipfile.ZipFile(path) as archive:
        _check_archive(archive)
        presentation = _read_xml(archive, 'ppt/presentation.xml')
        relations = _read_xml(archive, 'ppt/_rels/presentation.xml.rels')
        targets = {}
        for rel in relations:
            if rel.get('Type', '').endswith('/slide'):
                if rel.get('TargetMode') == 'External':
                    raise ValueError('external slides unsupported')
                target = rel.get('Target', '')
                if (not target or target.startswith('//') or '..' in target.split('/')
                        or any(char in target for char in ('\\', ':', '?', '#'))):
                    raise ValueError('invalid slide target')
                # OPC relationships may use a package-root URI (/ppt/...), not
                # an operating-system absolute path. Never resolve on disk.
                name = posixpath.normpath(target[1:] if target.startswith('/')
                                         else posixpath.join('ppt', target))
                if not name.startswith('ppt/slides/') or not name.endswith('.xml'):
                    raise ValueError('invalid slide target')
                targets[rel.get('Id')] = name
        slides = presentation.findall(f'{{{P}}}sldIdLst/{{{P}}}sldId')
        if not slides or len(slides) > MAX_PAGES:
            raise ValueError('invalid slide count')
        for index, slide in enumerate(slides, 1):
            root = _read_xml(archive, targets[slide.get(f'{{{R}}}id')])
            output.append(f'第{index}页')

            def walk(node):
                nonlocal found
                if node.tag == f'{{{A}}}tbl':
                    for row in node.findall(f'{{{A}}}tr'):
                        cells = [' / '.join(_drawing_paragraph(p) for p in cell.iter(f'{{{A}}}p'))
                                 for cell in row.findall(f'{{{A}}}tc')]
                        if any(cell.strip() for cell in cells):
                            output.append(' | '.join(cells))
                            found = True
                    return
                if node.tag == f'{{{A}}}p':
                    text = _drawing_paragraph(node)
                    if text.strip():
                        output.append(text)
                        found = True
                    return
                for child in node:
                    walk(child)

            walk(root)
    return 'ready' if found else 'needs_ocr'


def _inflate_diagram(payload):
    compressed = base64.b64decode(''.join(payload.split()), validate=True)
    decoder = zlib.decompressobj(wbits=-15)
    decoded = decoder.decompress(compressed, MAX_XML_BYTES + 1)
    if len(decoded) > MAX_XML_BYTES or decoder.unconsumed_tail or not decoder.eof or decoder.unused_data:
        raise ValueError('diagram decompression limit or invalid payload')
    decoded = urllib.parse.unquote_to_bytes(decoded.decode('utf-8'))
    return _xml(decoded), len(decoded)


def _drawio(path, output):
    root = _xml(path.read_bytes())
    if root.tag == 'mxGraphModel':
        diagrams = [(None, root)]
    elif root.tag == 'mxfile':
        diagrams = []
        expanded_bytes = 0
        for diagram in root.findall('diagram'):
            graph = diagram.find('mxGraphModel')
            if graph is None:
                payload = (diagram.text or '').strip()
                if not payload:
                    return 'archived'
                graph, size = _inflate_diagram(payload)
                expanded_bytes += size
                if expanded_bytes > MAX_ZIP_BYTES:
                    raise ValueError('combined diagram expansion limit exceeded')
            if graph.tag != 'mxGraphModel':
                return 'archived'
            diagrams.append((diagram.get('name'), graph))
    else:
        return 'archived'
    if len(diagrams) > MAX_PAGES:
        raise ValueError('too many diagrams')
    found = False
    for index, (name, graph) in enumerate(diagrams, 1):
        output.append(f'第{index}页' + (f'：{name}' if name else ''))
        for node in graph.iter():
            # Draw.io wraps custom business objects around an mxCell.
            if node.tag == 'mxCell':
                label = node.get('value', '')
            elif node.tag.lower() in {'object', 'userobject'}:
                label = node.get('label', '')
            else:
                continue
            text = _html_text(label)
            if text:
                output.append(text)
                found = True
    return 'ready' if found else 'archived'


def _pdf(path, output):
    from pypdf import filters
    names = ('ZLIB_MAX_OUTPUT_LENGTH', 'LZW_MAX_OUTPUT_LENGTH', 'RUN_LENGTH_MAX_OUTPUT_LENGTH',
             'JBIG2_MAX_OUTPUT_LENGTH', 'MAX_ARRAY_BASED_STREAM_OUTPUT_LENGTH', 'MAX_DECLARED_STREAM_LENGTH')
    original = {name: getattr(filters, name) for name in names if hasattr(filters, name)}
    try:
        # Bound internal font/XObject and object-stream expansion as well.
        for name in original:
            setattr(filters, name, MAX_PDF_STREAM_BYTES)
        return _pdf_text(path, output)
    finally:
        for name, value in original.items():
            setattr(filters, name, value)


def _pdf_text(path, output):
    from pypdf import PdfReader
    from pypdf.generic import ArrayObject, DictionaryObject, IndirectObject, StreamObject
    reader = PdfReader(path)
    if reader.is_encrypted or len(reader.pages) > MAX_PAGES:
        return 'failed'
    visited = set()
    stream_bytes = 0

    def validate_streams(value):
        nonlocal stream_bytes
        if isinstance(value, IndirectObject):
            key = ('ref', value.idnum, value.generation)
            if key in visited:
                return
            visited.add(key)
            value = value.get_object()
        key = id(value)
        if key in visited:
            return
        visited.add(key)
        if len(visited) > 50000:
            raise ValueError('too many PDF objects')
        if isinstance(value, StreamObject):
            if value.get('/Subtype') == '/Image':
                return  # Text extraction does not decode raster pixels.
            data = value.get_data()
            stream_bytes += len(data)
            if len(data) > MAX_PDF_STREAM_BYTES or stream_bytes > MAX_ZIP_BYTES:
                raise ValueError('PDF stream expansion limit exceeded')
        if isinstance(value, DictionaryObject):
            for key, child in value.items():
                if key not in {'/Parent', '/P'}:
                    validate_streams(child)
        elif isinstance(value, ArrayObject):
            for child in value:
                validate_streams(child)

    found = False
    for index, page in enumerate(reader.pages, 1):
        # Validate ahead of extraction: pypdf may otherwise log and skip a failed
        # XObject, leaving a misleading partial success.
        validate_streams(page.get('/Contents'))
        validate_streams(page.get('/Resources'))
        text = page.extract_text() or ''
        found = found or bool(text.strip())
        output.append(f'第{index}页\n' + text)
    return 'ready' if found else 'needs_ocr'


def _extract(path):
    ext = path.suffix.lower()
    if path.stat().st_size > (MAX_PDF_BYTES if ext == '.pdf' else MAX_FILE_BYTES):
        return '', 'failed'
    if ext in IMAGE_EXTENSIONS:
        return '', 'needs_ocr'
    if ext in {'.docx', '.xlsx'}:
        with zipfile.ZipFile(path) as archive:
            _check_archive(archive)
    output = _Text()
    status = 'ready'
    if ext in {'.md', '.txt'}:
        output.append(path.read_text(encoding='utf-8-sig'))
    elif ext in {'.html', '.htm'}:
        output.append(_html_text(path.read_text(encoding='utf-8-sig')))
        if not output.text().strip():
            status = 'archived'
    elif ext == '.pptx':
        status = _pptx(path, output)
    elif ext == '.drawio':
        status = _drawio(path, output)
    elif ext == '.pdf':
        status = _pdf(path, output)
    elif ext == '.docx':
        from docx import Document
        document = Document(path)
        for index, paragraph in enumerate(document.paragraphs, 1):
            if paragraph.text:
                output.append(f'段落{index}：{paragraph.text}')
        for index, table in enumerate(document.tables, 1):
            for row_index, row in enumerate(table.rows, 1):
                output.append(f'表{index} 行{row_index}：' + ' | '.join(c.text for c in row.cells))
    elif ext == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            cells = 0
            for sheet in workbook:
                output.append('工作表：' + sheet.title)
                for row in sheet.iter_rows():
                    cells += len(row)
                    if cells > 200000:
                        raise ValueError('too many cells')
                    values = [str(cell.coordinate) + '=' + str(cell.value) for cell in row if cell.value is not None]
                    if values:
                        output.append(' | '.join(values))
        finally:
            workbook.close()
    else:
        status = 'archived'
    return (output.text() if status == 'ready' else ''), status


def extract(path):
    """Return complete text and a parse status; any parse/limit failure is empty."""
    try:
        text, status = _extract(Path(path))

        def repair_surrogate(match):
            value = match.group()
            if len(value) == 2:
                return chr(0x10000 + ((ord(value[0]) - 0xD800) << 10)
                           + ord(value[1]) - 0xDC00)
            return '[无法辨认字符]'

        # PDF ToUnicode maps can yield UTF-16 code units instead of Unicode
        # scalars. Preserve valid pairs; mark unmatched units visibly rather
        # than guessing a missing glyph or passing invalid UTF-8 to storage.
        text = re.sub('[\ud800-\udbff][\udc00-\udfff]|[\ud800-\udfff]', repair_surrogate, text)
        if len(text) > MAX_CHARS:
            raise ValueError('normalized document too long')
        return text, status
    except Exception:
        return '', 'failed'


if __name__ == '__main__':
    text, status = extract(sys.argv[1])
    print(json.dumps({'text': text, 'status': status}, ensure_ascii=True))
