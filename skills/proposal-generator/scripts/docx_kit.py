"""方案文档排版工具（规范说明见 references/04_排版规范.md）。

用 d.mute = True/False 包住要从精简版（如汇报版）中去掉的内容，同一份正文可生成两个版本。

字体规范（商务方案通行格式）：
  标题      黑体：一级 三号(16pt)、二级 四号(14pt)、三级 小四(12pt)
  正文      宋体 小四(12pt)，1.5 倍行距，首行缩进 2 字符
  表格      表头 黑体 五号(10.5pt)；内容 宋体 五号(10.5pt)
  图题表题  黑体 五号(10.5pt) 居中；表题在表上，图题在图下，自动编号
  页眉页脚  宋体 小五(9pt)
  西文与数字 Times New Roman
"""
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

SONG, HEI, LATIN = "宋体", "黑体", "Times New Roman"
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x59, 0x59, 0x59)
LINE = "A6B1BD"
HEAD_FILL = "E6ECF2"
KEY_FILL = "F2F5F8"
GROUP_FILL = "F2F5F8"
CONTENT_WIDTH_CM = 16.0

SIZE = {"一号": 26, "小一": 24, "二号": 22, "小二": 18, "三号": 16, "四号": 14, "小四": 12, "五号": 10.5, "小五": 9}


def _fonts(target, east, size=None, bold=None, color=None, latin=LATIN):
    """target 可以是 run 或 style。"""
    f = target.font
    f.name = latin
    rpr = target.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    # 内置样式带主题字体属性，优先级高于具体字体，必须去掉
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rf.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(attr), latin)
    rf.set(qn("w:eastAsia"), east)
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color is not None:
        f.color.rgb = color


def _pf(p, before=0, after=0, line=1.5, first_indent=False, align=None, keep_next=False):
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line
    if first_indent:
        ind = OxmlElement("w:ind")
        ind.set(qn("w:firstLineChars"), "200")
        p._p.get_or_add_pPr().append(ind)
    if align is not None:
        p.alignment = align
    if keep_next:
        pf.keep_with_next = True


def _rich(p, text, east=SONG, size=SIZE["小四"], color=BLACK, bold_all=False):
    """支持 **强调** 标记；强调片段用黑体。"""
    for i, part in enumerate(re.split(r"\*\*", text)):
        if not part:
            continue
        r = p.add_run(part)
        emph = bold_all or i % 2 == 1
        _fonts(r, HEI if emph else east, size, bold=False, color=color)


def _field(p, instr, size=SIZE["小五"]):
    r = p.add_run()
    _fonts(r, SONG, size, color=GRAY)
    for kind, val in (("begin", None), ("instr", instr), ("separate", None), ("text", "1"), ("end", None)):
        if kind in ("begin", "separate", "end"):
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = val
        else:
            el = OxmlElement("w:t")
            el.text = val
        r._r.append(el)


def _shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _table_borders(table):
    b = OxmlElement("w:tblBorders")
    for side, sz, color in (("top", 12, "3E5C76"), ("bottom", 12, "3E5C76"), ("left", 4, LINE),
                            ("right", 4, LINE), ("insideH", 4, LINE), ("insideV", 4, LINE)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), color)
        b.append(el)
    table._tbl.tblPr.append(b)


def _cell_margins(table, top=60, bottom=60, left=110, right=110):
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    table._tbl.tblPr.append(mar)


def _fix_grid(table, widths_cm):
    """写入 tblGrid、表宽和固定布局，避免 Word 均分列宽。"""
    tbl = table._tbl
    twips = [int(w / 2.54 * 1440) for w in widths_cm]
    grid = tbl.tblGrid
    for gc in list(grid):
        grid.remove(gc)
    for w in twips:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(w))
        grid.append(gc)
    tblpr = tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        old = tblpr.find(qn(tag))
        if old is not None:
            tblpr.remove(old)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(sum(twips)))
    tw.set(qn("w:type"), "dxa")
    tblpr.append(tw)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblpr.append(lay)


def _row_flag(row, tag):
    el = OxmlElement(f"w:{tag}")
    if tag == "tblHeader":
        el.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(el)


def _cell_text(cell, text, east=SONG, size=SIZE["五号"], align=None):
    for j, line in enumerate(str(text).split("\n")):
        p = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
        _pf(p, line=1.25, align=align)
        _rich(p, line, east=east, size=size)


class Doc:
    def __init__(self, header_text):
        self.d = Document()
        self.header_text = header_text
        self.n_table = 0
        self.n_fig = 0
        self.mute = False  # True 时跳过输出，用于从同一份正文生成精简版
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.5)
        sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.3)
        sec.header_distance, sec.footer_distance = Cm(1.5), Cm(1.4)
        st = self.d.styles
        _fonts(st["Normal"], SONG, SIZE["小四"], color=BLACK)
        for name, size, before, after in (("Heading 1", SIZE["三号"], 18, 10), ("Heading 2", SIZE["四号"], 12, 6),
                                          ("Heading 3", SIZE["小四"], 8, 4)):
            s = st[name]
            _fonts(s, HEI, size, bold=False, color=BLACK)
            pf = s.paragraph_format
            pf.space_before, pf.space_after = Pt(before), Pt(after)
            pf.line_spacing = 1.3
            pf.keep_with_next = True
        _fonts(st["List Bullet"], SONG, SIZE["小四"], color=BLACK)
        for name, east, indent in (("toc 1", HEI, 0), ("toc 2", SONG, 0.74)):
            try:
                s = st[name]
            except KeyError:
                s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
                # 按 Word 内置目录样式的 ID 创建（TOC1/TOC2），否则 Word 会忽略并改用自带样式
                s.element.set(qn("w:styleId"), name.replace("toc ", "TOC"))
                s.element.attrib.pop(qn("w:customStyle"), None)
            _fonts(s, east, SIZE["小四"], color=BLACK)
            s.paragraph_format.left_indent = Cm(indent)
            # 目录要留出页尾余量：目录正好排满一页时，域结束标记会被挤到下一页形成空白页
            s.paragraph_format.space_after = Pt(1)
            s.paragraph_format.line_spacing = 1.1

    # ---------- 封面与目录 ----------
    def cover(self, org, title_lines, tagline, meta_lines):
        for _ in range(7):
            _pf(self.d.add_paragraph(), line=1.0)
        p = self.d.add_paragraph()
        _pf(p, after=20, line=1.2, align=WD_ALIGN_PARAGRAPH.CENTER)
        _fonts(p.add_run(org), HEI, SIZE["小二"])
        for t in title_lines:
            p = self.d.add_paragraph()
            _pf(p, after=8, line=1.2, align=WD_ALIGN_PARAGRAPH.CENTER)
            _fonts(p.add_run(t), HEI, SIZE["一号"])
        p = self.d.add_paragraph()
        _pf(p, before=16, after=160, line=1.2, align=WD_ALIGN_PARAGRAPH.CENTER)
        self._rule(p, top=True)
        _fonts(p.add_run(tagline), SONG, SIZE["四号"], color=GRAY)
        for line in meta_lines:
            p = self.d.add_paragraph()
            _pf(p, after=8, line=1.2, align=WD_ALIGN_PARAGRAPH.CENTER)
            _fonts(p.add_run(line), SONG, SIZE["四号"])
        new = self.d.add_section(WD_SECTION.NEW_PAGE)
        new.header.is_linked_to_previous = False
        new.footer.is_linked_to_previous = False
        hp = new.header.paragraphs[0]
        _pf(hp, line=1.0, align=WD_ALIGN_PARAGRAPH.RIGHT)
        _fonts(hp.add_run(self.header_text), SONG, SIZE["小五"], color=GRAY)
        self._rule(hp, top=False)
        fp = new.footer.paragraphs[0]
        _pf(fp, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        _field(fp, "PAGE")

    @staticmethod
    def _rule(p, top):
        pbdr = OxmlElement("w:pBdr")
        el = OxmlElement("w:top" if top else "w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "10" if top else "4"), ("w:color", LINE)):
            el.set(qn(k), v)
        pbdr.append(el)
        p._p.get_or_add_pPr().append(pbdr)

    def toc(self):
        p = self.d.add_paragraph()
        _pf(p, after=18, line=1.2, align=WD_ALIGN_PARAGRAPH.CENTER)
        _fonts(p.add_run("目　　录"), HEI, SIZE["三号"])
        p = self.d.add_paragraph()
        _field(p, 'TOC \\o "1-2" \\h \\z \\u', size=SIZE["小四"])
        # 不插分页符：目录满一页时分页符会被挤成空白页；改由下一个一级标题“段前分页”
        self._break_before_next_h1 = True

    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---------- 文本 ----------
    def h1(self, text):
        if self.mute:
            return
        h = self.d.add_heading(text, level=1)
        if getattr(self, "_break_before_next_h1", False):
            h.paragraph_format.page_break_before = True
            self._break_before_next_h1 = False

    def h2(self, text):
        if self.mute:
            return
        self.d.add_heading(text, level=2)

    def h3(self, text):
        if self.mute:
            return
        self.d.add_heading(text, level=3)

    def p(self, text):
        if self.mute:
            return
        p = self.d.add_paragraph()
        _pf(p, first_indent=True)
        _rich(p, text)
        return p

    def note(self, text):
        if self.mute:
            return
        p = self.d.add_paragraph()
        _pf(p, before=0, after=6, line=1.3)
        _rich(p, "注：" + text, size=SIZE["五号"], color=GRAY)

    def bullets(self, items):
        if self.mute:
            return
        for it in items:
            p = self.d.add_paragraph(style="List Bullet")
            _pf(p)
            _rich(p, it)

    # ---------- 表格 ----------
    def table(self, title, headers, rows, widths, center_cols=(), key_col=True):
        """rows 中 ('__group__', 文本) 渲染为通栏分组行；key_col=True 时首列浅底黑体居中。"""
        if self.mute:
            return
        assert abs(sum(widths) - CONTENT_WIDTH_CM) < 0.05, (title, sum(widths))
        self._table_title(title)
        t = self.d.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        _table_borders(t)
        _cell_margins(t)
        hdr = t.rows[0]
        _row_flag(hdr, "tblHeader")
        for i, h in enumerate(headers):
            c = hdr.cells[i]
            _shade(c, HEAD_FILL)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _cell_text(c, h, east=HEI, align=WD_ALIGN_PARAGRAPH.CENTER)
        for row in rows:
            r = t.add_row()
            _row_flag(r, "cantSplit")
            if isinstance(row, tuple) and row and row[0] == "__group__":
                m = r.cells[0].merge(r.cells[-1])
                _shade(m, GROUP_FILL)
                _cell_text(m, row[1], east=HEI)
                continue
            for i, val in enumerate(row):
                c = r.cells[i]
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                is_key = key_col and i == 0
                if is_key:
                    _shade(c, KEY_FILL)
                _cell_text(c, val, east=HEI if is_key else SONG,
                           align=WD_ALIGN_PARAGRAPH.CENTER if (is_key or i in center_cols) else None)
        for row in t.rows:
            for i, c in enumerate(row.cells):
                if i < len(widths):
                    c.width = Cm(widths[i])
        # 只保证表头与第一行数据不分离；表格可跨页续排（表头重复），避免整表跳页留下大片空白
        keep_rows = t.rows[:1]
        for row in keep_rows:
            for c in row.cells:
                for para in c.paragraphs:
                    para.paragraph_format.keep_with_next = True
        _fix_grid(t, widths)
        self._gap()

    def kv(self, title, rows, widths=(3.0, 13.0)):
        """两列“项目—内容”表，用于把要点式文字整理成表。"""
        self.table(title, ["项目", "内容"], rows, list(widths))

    def _table_title(self, title):
        prev = self.d.paragraphs[-1]
        if prev.text.strip() and not prev.text.startswith("图 "):  # 引导语与表格同页；图题不绑定，避免大块留白
            prev.paragraph_format.keep_with_next = True
        self.n_table += 1
        p = self.d.add_paragraph()
        _pf(p, before=6, after=4, line=1.2, align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
        _fonts(p.add_run(f"表 {self.n_table}  {title}"), HEI, SIZE["五号"])

    # ---------- 图 ----------
    def image(self, path, width_cm, caption):
        if self.mute:
            return
        self.d.paragraphs[-1].paragraph_format.keep_with_next = True
        p = self.d.add_paragraph()
        _pf(p, before=6, after=2, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        self.n_fig += 1
        c = self.d.add_paragraph()
        _pf(c, after=10, line=1.2, align=WD_ALIGN_PARAGRAPH.CENTER)
        _fonts(c.add_run(f"图 {self.n_fig}  {caption}"), HEI, SIZE["五号"])

    def _gap(self):
        _pf(self.d.add_paragraph(), line=0.6)

    def save(self, path):
        # 去掉文末空段落，避免多出空白页
        while self.d.paragraphs and not self.d.paragraphs[-1].text.strip() \
                and not self.d.paragraphs[-1]._p.xpath(".//w:drawing | .//w:br"):
            last = self.d.paragraphs[-1]._p
            if last.getnext() is not None and last.getnext().tag != qn("w:sectPr"):
                break
            last.getparent().remove(last)
        # Word 要求表格后必须有段落：用 1 磅高的极小段落收尾，避免被挤到新页
        tail = self.d.add_paragraph()
        tail.paragraph_format.space_before = tail.paragraph_format.space_after = Pt(0)
        tail.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        tail.paragraph_format.line_spacing = Pt(1)
        rpr = tail._p.get_or_add_pPr()
        mark = OxmlElement("w:rPr")
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "2")
        mark.append(sz)
        rpr.append(mark)
        self.d.save(path)
