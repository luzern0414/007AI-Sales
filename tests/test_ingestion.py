import io
from docx import Document
from openpyxl import Workbook
from tests.test_platform import clients,create_project


def test_office_upload_preserves_searchable_sources(clients):
    a,_=clients
    p=create_project(a)
    document=Document();document.add_paragraph('客户要求权限审批')
    doc=io.BytesIO();document.save(doc)
    sheet=Workbook();sheet.active.title='权限规则';sheet.active.append(['角色','权限']);sheet.active.append(['财务主管','审批报销'])
    book=io.BytesIO();sheet.save(book)
    for name,content in [('meeting.docx',doc.getvalue()),('rules.xlsx',book.getvalue())]:
        r=a.post(f'/api/projects/{p["id"]}/files',files={'file':(name,content)})
        assert r.status_code==201 and r.json()['parse_status']=='ready',r.text
    results=a.get(f'/api/projects/{p["id"]}/search',params={'q':'审批'}).json()
    assert len(results)==2
    assert any('权限规则' in r['snippet'] for r in results)


def test_same_filename_is_versioned_without_overwrite(clients):
    a,_=clients
    p=create_project(a)
    url=f'/api/projects/{p["id"]}/files'
    first=a.post(url,files={'file':('scope.md',b'original scope')}).json()
    second=a.post(url,files={'file':('scope.md',b'new scope')}).json()
    assert first['id']!=second['id']
    assert second['previous_version_id']==first['id'] and second['file_version']==2
    assert a.get(f'/api/files/{first["id"]}/download').content==b'original scope'
