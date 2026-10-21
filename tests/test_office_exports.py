import io
import pytest
from docx import Document
from openpyxl import load_workbook


def test_word_draft_preserves_text_and_fixed_table_layout():
    from server.office_exports import word_draft
    content='# 审批项目\n\n## 一、需求范围\n\n下表列出甲方待确认需求。\n\n| 编号 | 内容 |\n| --- | --- |\n| REQ-01 | 审批权限待确认 |\n'
    doc=Document(io.BytesIO(word_draft('审批实施方案','隔离测试项目',content)))
    assert any('讨论稿' in p.text for p in doc.paragraphs)
    assert doc.tables[0].cell(1,1).text=='审批权限待确认'
    assert not doc.tables[0].autofit
    section=doc.sections[0]
    assert abs(section.page_width.cm-21)<0.01
    assert abs(section.left_margin.cm-2.5)<0.01
    xml=doc.element.xml
    assert 'tblHeader' in xml and 'cantSplit' in xml
    assert 'SimSun' in xml and 'SimHei' in xml


def test_checklist_workbook_preserves_fields_without_formula_execution():
    from server.office_exports import checklist_book
    md='| 编号 | 分类 | 事项 | 需要准备的内容 | 用途 | 形式 | 交付格式 | 参考模板 | 必需性 | 建议责任方 | 状态 | 甲方填写 / 备注 |\n| --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- |\n| C-01 | 权限 | 角色确认 | =HYPERLINK("https://invalid.test") | 配置权限 | 书面 | Excel | | 必需 | 待确认 | 待提供 | |'
    wb=load_workbook(io.BytesIO(checklist_book('资料确认清单',md,'client_checklist')))
    ws=wb.worksheets[0]
    assert ws.cell(4,1).value=='C-01'
    assert ws.cell(4,4).data_type=='s'
    assert ws.freeze_panes=='D4' and ws.auto_filter.ref=='A3:L4'
    assert '一期' not in ' '.join(wb.sheetnames)
    assert ws.cell(3,12).value=='甲方填写 / 备注'
    assert len(ws.data_validations.dataValidation)==1


def test_checklist_rejects_missing_structure_instead_of_inventing_rows():
    from server.office_exports import checklist_book
    with pytest.raises(ValueError,match='表格'):
        checklist_book('需求清单','# 需求清单\n\n权限待确认','requirements')


def test_office_download_is_owned_and_validates_format(clients):
    from tests.test_platform import create_project
    from tests.test_recordings import new_meeting
    from server.db import GeneratedDocument
    a,b=clients;p=create_project(a);m=new_meeting(a,p)
    with a.app.state.sessions() as db:
        d=GeneratedDocument(project_id=p['id'],meeting_id=m['id'],kind='proposal',title='审批实施方案',markdown='# 审批实施方案\n\n## 一、需求\n\n权限待确认',source_snapshot='[]',source_revision=1)
        db.add(d);db.commit();did=d.id
    url=f'/api/documents/{did}/download'
    assert b.get(url+'?format=docx').status_code==401
    r=a.get(url+'?format=docx')
    assert r.status_code==200
    assert r.headers['content-type']=='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    assert Document(io.BytesIO(r.content)).paragraphs
    assert a.get(url+'?format=xlsx').status_code==422


from tests.test_platform import clients
